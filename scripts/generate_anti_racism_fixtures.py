"""Generate fictional Anti-Racism Call cases with repeated activity records.

The supplied PDF is never read by this generator or copied into its output.
"""
import argparse
import hashlib
import json
import shutil
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader

from generate_fixtures import convert

ROOT = Path(__file__).resolve().parents[1]
VARIANTS = ["continuation_no_header", "repeated_header", "split_activity_row",
            "extra_rows", "ambiguous_selections", "blank_spacer_rows",
            "missing_counts", "reordered_columns"]


def answer(text, normalized=None, reasons=()):
    return {"answer_text": text, "normalized_answer": text if normalized is None else normalized,
            "review_required": bool(reasons), "review_reasons": list(reasons)}


def ordinary_answers(schema, n):
    values={}
    for f in schema["questions"]:
        typ=f["response_type"]
        text={"Text":"Fictional test response", "Telephone":"202-555-0100",
              "Email":f"arc{n:03d}@example.invalid", "Date":"2026-06-20",
              "Currency":"12500.00", "Number":"12", "Yes / No":"Yes",
              "Text / links":"https://example.invalid/project",
              "Written response":"Fictional participants discussed inclusive practices and shared feedback.",
              "Signature":"Alex Example (synthetic typed signature)"}[typ]
        values[f["field_id"]]=text
    values.update({"ARC-A01":f"Fictional Community Project {n:03d}", "ARC-A02":"Alex", "ARC-A03":"Example",
                   "ARC-A04":"Coordinator", "ARC-A06":"N/A", "ARC-B01":f"Synthetic Inclusion Project {n:03d}",
                   "ARC-B02":"2026-05-01", "ARC-B03":"2026-06-13", "ARC-B05":"10000.00",
                   "ARC-C01":"No", "ARC-C02":"Yes", "ARC-C03":"Ottawa, Ontario; online. No change of location.",
                   "ARC-D09":"No", "ARC-D09a":"N/A", "ARC-S02":"Alex Example"})
    return values


def rows_for(variant):
    rows=[]
    for i in range(1,8 if variant=="extra_rows" else 5):
        description=f"Activity {i}: Fictional inclusion workshop. Participants discussed workplace access and shared examples."
        if variant=="extra_rows": description+=f"\nTopic {i}: access\nTopic {i}: participation"
        rows.append({"activity_index":i,"description":answer(description),
                     "completed":answer("Yes [X] / No [ ]","Yes"),
                     "participant_count":answer(str(i*10),i*10),"source_segments":[],
                     "review_required":False,"review_reasons":[]})
    if variant=="ambiguous_selections":
        for row,raw,reason in zip(rows[:3],["Yes [X] / No [X]","Yes [ ] / No [ ]","Yes / No [ ]"],
                                  ["both_options_selected","no_option_selected","selection_symbol_missing"]):
            row["completed"]=answer(raw,reasons=[reason]); row["completed"]["normalized_answer"]=None
        rows[3]["completed"]=answer("Yes [ ] / No [X]","No")
    if variant=="missing_counts":
        for row,raw,val,reasons in zip(rows,["0","","N/A","about 20"],[0,None,None,20],
                                       [[],["participant_count_blank"],[],["approximate_count"]]):
            row["participant_count"]=answer(raw,reasons=reasons); row["participant_count"]["normalized_answer"]=val
    for row in rows:
        row["review_reasons"]=list(dict.fromkeys(reason for key in ["description","completed","participant_count"] for reason in row[key]["review_reasons"]))
        row["review_required"]=bool(row["review_reasons"])
    return rows


def write_case(path, schema, variant, n):
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Inches(.65)
    doc.styles["Normal"].font.name="Arial"; doc.styles["Normal"].font.size=Pt(10)
    for name in ["Title","Heading 1","Heading 2"]: doc.styles[name].font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph("MARP Final Activity Report Anti-Racism Call",style="Title")
    doc.add_paragraph(f"SYNTHETIC TEST FORM ONLY. Case {n:02d}. Layout: {variant}.")
    sec.header.paragraphs[0].text="FICTIONAL DATA ONLY"
    values=ordinary_answers(schema,n); rows=rows_for(variant); expected={}; source_rows=[]
    table_index=0

    def response_field(f):
        nonlocal table_index
        table=doc.add_table(rows=2,cols=1); table.style="Table Grid"; table_index+=1
        table.cell(0,0).text=f["question"]
        for run in table.cell(0,0).paragraphs[0].runs: run.bold=True
        value=values[f["field_id"]]
        table.cell(1,0).text=("Yes [X] / No [ ]" if value=="Yes" else "Yes [ ] / No [X]") if f["response_type"]=="Yes / No" else value
        reasons=["source_question_incomplete"] if f["wording_status"]=="incomplete_in_source" else []
        expected[f["field_id"]]=answer(value,reasons=reasons)
        expected[f["field_id"]]["source_location"]={"table_index":table_index,"row":2}
        doc.add_paragraph().paragraph_format.space_after=Pt(0)

    for section in ["A","B"]:
        doc.add_heading(schema["sections"][section],level=1)
        for f in schema["questions"]:
            if f["section"]==section: response_field(f)
    doc.add_page_break(); doc.add_heading("C. Project Activities",level=1)
    doc.add_paragraph("List and describe your project’s activities. Record completion and participants for each activity.")
    order=["description","completed","participant_count"]
    if variant=="reordered_columns": order=["participant_count","description","completed"]
    labels={"description":"Activity", "completed":"Completed", "participant_count":"Number of participants per activity (if applicable)"}

    def activity_table(header):
        nonlocal table_index
        t=doc.add_table(rows=0,cols=3); t.style="Table Grid"; t.autofit=False; table_index+=1
        widths={"description":3.6,"completed":1.2,"participant_count":1.7}
        for i,key in enumerate(order): t.columns[i].width=Inches(widths[key])
        if header:
            cells=t.add_row().cells
            for i,key in enumerate(order): cells[i].text=labels[key]
        return t

    table=activity_table(True)
    for index,row in enumerate(rows):
        if index==2:
            doc.add_page_break()
            table=activity_table(variant=="repeated_header")
        if variant=="blank_spacer_rows" and index==1:
            table.add_row()
            source_rows.append({"table_index":table_index,"row":len(table.rows),"kind":"empty_spacer","activity_index":None})
        cells=table.add_row().cells
        for i,key in enumerate(order): cells[i].text=row[key]["answer_text"]
        row["source_segments"].append({"table_index":table_index,"row":len(table.rows)})
        source_rows.append({"table_index":table_index,"row":len(table.rows),"kind":"activity","activity_index":row["activity_index"]})
        if variant=="split_activity_row" and index==1:
            cells[order.index("description")].text=f"Activity {index+1}: Fictional inclusion workshop."
            doc.add_page_break(); table=activity_table(False); cells=table.add_row().cells
            cells[order.index("description")].text="Participants discussed workplace access and shared examples."
            row["source_segments"].append({"table_index":table_index,"row":len(table.rows)})
            source_rows.append({"table_index":table_index,"row":len(table.rows),"kind":"continuation","activity_index":row["activity_index"]})
            row["review_required"]=True; row["review_reasons"].append("activity_row_continuation")
    doc.add_page_break()
    for section in ["C","D","E","F","S"]:
        doc.add_heading(schema["sections"][section],level=1)
        for f in schema["questions"]:
            if f["section"]==section: response_field(f)
    doc.save(path)
    return {"responses":expected,"activities":rows,"physical_activity_rows":source_rows,
            "document_review_required":True,"document_review_reasons":["source_question_number_gap","source_unlabelled_response_boxes"]}


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--output",type=Path,default=ROOT/"fixtures/anti_racism_call")
    parser.add_argument("--soffice",default=shutil.which("soffice"))
    args=parser.parse_args()
    if not args.soffice: parser.error("LibreOffice soffice is required")
    out=args.output.resolve()
    if out.exists() and any(out.iterdir()): parser.error("Output directory must be empty")
    (out/"docx").mkdir(parents=True,exist_ok=True); (out/"pdf").mkdir()
    schema=json.loads((ROOT/"schemas/anti_racism_call.json").read_text())
    records=[]; paths=[]
    for n,variant in enumerate(VARIANTS,1):
        stem=f"arc_{n:02d}_{variant}"; path=out/"docx"/(stem+".docx")
        truth=write_case(path,schema,variant,n); paths.append(path)
        for ext in ["docx","pdf"]:
            records.append({"file":f"{ext}/{stem}.{ext}","case_id":stem,"template_id":schema["template_id"],"variant":variant,"expected":truth})
    convert(paths,out/"pdf",args.soffice)
    for r in records:
        p=out/r["file"]; r["sha256"]=hashlib.sha256(p.read_bytes()).hexdigest()
        if p.suffix==".pdf": r["pages"]=len(PdfReader(p).pages)
    (out/"expected.jsonl").write_text("\n".join(json.dumps(r,ensure_ascii=False) for r in records)+"\n",encoding="utf-8")
    (out/"manifest.json").write_text(json.dumps({"template_id":schema["template_id"],"files":len(records),"docx":8,"pdf":8,"variants":VARIANTS,"synthetic":True},indent=2)+"\n")
    print(f"Generated {len(records)} Anti-Racism Call files")


if __name__=="__main__": main()
