"""Generate fictional MARP forms and independent expected-response records.

Requires LibreOffice and Poppler (soffice, pdftoppm) on PATH.
Run from any directory: python scripts/generate_fixtures.py
"""
import argparse
import hashlib
import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageFilter
from pypdf import PdfReader
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
VARIANTS = ["ascii_x", "unicode_checks", "native_controls", "both_selected",
            "neither_selected", "reversed_options", "typed_answers", "circled_choice",
            "crossed_out", "extra_rows", "bulleted_lists", "page_continuation",
            "missing_question", "reordered_sections", "numeric_edge_cases",
            "contradictory_followup", "header_words_in_answer", "merged_cells",
            "long_answers", "accents_and_mixed_fonts"]
YESNO = {"C02", "D02", "G01"}


def base_answers(n):
    return dict(zip([q[0] for q in json.loads((ROOT / "questions.json").read_text())], [
        f"Fictional Community Group {n:04d}", "Alex", "Example", "Event coordinator",
        "202-555-0100", "N/A", f"fixture{n:04d}@example.invalid",
        f"Synthetic Culture Day {n:04d}", "2026-06-12", "2026-06-13", "Ottawa, Ontario",
        "https://example.invalid/event", "Workshop: 30 participants\nMusic session: 45 participants\nDiscussion: 25 participants",
        "Yes", "The workshop moved indoors because of rain.",
        "Participants learned about cultural traditions and discussed inclusive community activities.",
        "Yes", "Fictional River Group: hosted music\nFictional Hill Group: hosted discussions", "N/A", "100",
        "3 occurrences", "Participants reported more confidence meeting new neighbours.",
        "Volunteers, accessible venues, and advance communication.", "Transport was limited. Next time we would arrange a shuttle.",
        "Fictional Library: venue\nFictional Arts Club: volunteers\nFictional Cafe: refreshments",
        "Posters, email invitations, and community social media.", "Yes",
        "Support was acknowledged on posters and at the closing event.", "N/A",
        "Invitations and announcements were provided in English and French.",
        "Alex Example (synthetic typed signature)", "Alex Example", "2026-06-20"]))


def checkbox(paragraph, checked):
    control = OxmlElement("w:sdt")
    props = OxmlElement("w:sdtPr")
    box = OxmlElement("w14:checkbox")
    state = OxmlElement("w14:checked"); state.set(qn("w14:val"), "1" if checked else "0")
    box.append(state)
    for tag, val in [("checkedState", "2612"), ("uncheckedState", "2610")]:
        el = OxmlElement("w14:" + tag); el.set(qn("w14:val"), val); el.set(qn("w14:font"), "DejaVu Sans"); box.append(el)
    props.append(box); control.append(props)
    content = OxmlElement("w:sdtContent"); run = OxmlElement("w:r"); text = OxmlElement("w:t")
    text.text = "☒" if checked else "☐"; run.append(text); content.append(run); control.append(content)
    paragraph._p.append(control)


def selection(paragraph, variant, answer):
    if variant == "native_controls":
        checkbox(paragraph, answer == "Yes"); paragraph.add_run(" Yes   ")
        checkbox(paragraph, answer == "No"); paragraph.add_run(" No")
        return
    if variant == "crossed_out":
        paragraph.add_run("[X] Yes").font.strike = True
        paragraph.add_run("   [X] No")
        return
    if variant == "circled_choice":
        paragraph.add_run("( Yes )     No" if answer == "Yes" else "Yes     ( No )")
        return
    if variant == "typed_answers":
        paragraph.add_run(answer); return
    marks = {"Yes": answer == "Yes", "No": answer == "No"}
    if variant == "both_selected": marks = {"Yes": True, "No": True}
    if variant == "neither_selected": marks = {"Yes": False, "No": False}
    options = ["No", "Yes"] if variant == "reversed_options" else ["Yes", "No"]
    for option in options:
        symbol = ("☒" if marks[option] else "☐") if variant == "unicode_checks" else ("[X]" if marks[option] else "[ ]")
        paragraph.add_run(f"{symbol} {option}    ")


def make_docx(path, variant, n, questions):
    answers = base_answers(n)
    # Alternate clear Yes and No cases, with consistent conditional responses.
    if n % 2 == 0 or variant == "crossed_out":
        answers.update(C02="No", C03="N/A", D02="No", D03="N/A", D04="Fictional River Group hosted the event.", G01="No", G02="N/A", G03="The acknowledgement was accidentally omitted.")
    if variant == "contradictory_followup": answers.update(C02="No", C03="Yes, the location changed to the library.")
    if variant == "numeric_edge_cases": answers.update(D05="0", B02="06/07/2026", B03="2026-02-30", A06="", D06="1,000", B05="N/A")
    if variant == "accents_and_mixed_fonts": answers.update(A02="Zoé", A03="Exemple", B04="Montréal, Québec", D08="Bénévoles; échanges culturels; accueil bilingue.")
    if variant == "header_words_in_answer": answers["D09"] = 'Our poster used the words "E. Partnerships" and "Name of Event:". These are part of this answer, not new fields.'
    if variant in {"long_answers", "page_continuation"}:
        answers["C01"] = "\n".join(f"Activity {i}: Synthetic participants discussed inclusion, shared stories, and practised listening. Session {i} had {i+10} participants." for i in range(1, 19))
    omitted = {"D07"} if variant == "missing_question" else set()
    doc = Document()
    sec = doc.sections[0]; sec.top_margin = sec.bottom_margin = Inches(.6)
    normal = doc.styles["Normal"]; normal.font.name = "DejaVu Sans"; normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    doc.add_heading("Synthetic MARP final activity report", 0)
    doc.add_paragraph(f"FICTIONAL TEST DATA ONLY | Case {n:04d} | {variant}")
    doc.sections[0].header.paragraphs[0].text = "SYNTHETIC TEST FORM — no real respondent data"
    order = list(questions)
    if variant == "reordered_sections": order.sort(key=lambda q: ("BACFEDGS".index(q[0][0]), q[0]))
    last_section = None
    expected = {}
    for idx, (fid, section, prompt, typ, condition) in enumerate(order):
        reasons = []
        raw = answers[fid]
        normalized = raw
        if fid in YESNO and variant in {"both_selected", "neither_selected", "circled_choice", "crossed_out"}:
            reasons = [{"both_selected":"both_options_selected", "neither_selected":"no_option_selected", "circled_choice":"nonstandard_selection_mark", "crossed_out":"corrected_selection"}[variant]]
            if variant in {"both_selected", "neither_selected"}: normalized = None
        if variant == "numeric_edge_cases" and fid in {"B02", "B03"}: reasons = ["ambiguous_date" if fid == "B02" else "invalid_date"]
        if variant == "contradictory_followup" and fid in {"C02", "C03"}: reasons = ["contradictory_followup"]
        if fid in omitted: reasons = ["question_missing"]; normalized = None; raw = None
        expected[fid] = {"answer_text": raw, "normalized_answer": normalized,
                         "review_required": bool(reasons), "review_reasons": reasons,
                         "presence": "missing" if fid in omitted else ("blank" if raw == "" else "present")}
        if fid in YESNO:
            expected[fid]["selected_options"] = (["Yes", "No"] if variant == "both_selected" else [] if variant == "neither_selected" else [answers[fid]])
            if variant in {"both_selected", "neither_selected"}: expected[fid]["answer_text"] = "[X] Yes [X] No" if variant == "both_selected" else "[ ] Yes [ ] No"
        if fid in omitted: continue
        if section != last_section:
            doc.add_heading(section, level=1); last_section = section
        table = doc.add_table(rows=1, cols=2 if variant == "merged_cells" else 1)
        table.style = "Table Grid"
        cell = table.cell(0,0)
        if variant == "merged_cells": cell = cell.merge(table.cell(0,1))
        cell.paragraphs[0].add_run(prompt).bold = True
        if variant == "accents_and_mixed_fonts":
            for run in cell.paragraphs[0].runs: run.font.name = "Liberation Serif" if idx % 2 else "DejaVu Sans"; run.font.size = Pt(10 + idx % 3)
        if fid in YESNO: selection(table.add_row().cells[0].paragraphs[0], variant, answers[fid])
        elif variant == "page_continuation" and fid == "C01":
            lines = raw.splitlines()
            table.add_row().cells[0].text = "\n".join(lines[:8])
            doc.add_page_break(); doc.add_paragraph("Activities continued")
            continuation = doc.add_table(rows=1, cols=1); continuation.style = "Table Grid"
            continuation.cell(0,0).text = "\n".join(lines[8:])
        elif variant in {"extra_rows", "merged_cells"} and "\n" in raw:
            for item in raw.splitlines():
                cells = table.add_row().cells
                if variant == "merged_cells" and ":" in item:
                    cells[0].text, cells[1].text = item.split(":",1)
                else: cells[0].text = item
        elif variant == "bulleted_lists" and "\n" in raw:
            answer_cell = table.add_row().cells[0]
            for k,item in enumerate(raw.splitlines()):
                p = answer_cell.paragraphs[0] if k == 0 else answer_cell.add_paragraph()
                p.style = "List Bullet"; p.add_run(item)
        else: table.add_row().cells[0].text = raw
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    doc.save(path)
    return expected


def convert(docx_paths, pdf_dir, soffice):
    with tempfile.TemporaryDirectory(prefix="marp-lo-") as tmp:
        # Bound command lengths for optional thousand-form batches.
        for start in range(0,len(docx_paths),25):
            subprocess.run([soffice, "-env:UserInstallation=" + Path(tmp).as_uri(), "--headless", "--convert-to", "pdf", "--outdir", str(pdf_dir), *map(str,docx_paths[start:start+25])], check=True, stdout=subprocess.DEVNULL, timeout=300)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=ROOT / "fixtures" / "synthetic")
    parser.add_argument("--count", type=int, default=20, help="Number of DOCX cases; each also has a converted PDF")
    parser.add_argument("--soffice", default=shutil.which("soffice"))
    parser.add_argument("--pdftoppm", default=shutil.which("pdftoppm"))
    args = parser.parse_args()
    if args.count < 20: parser.error("Use at least 20 to cover every scenario")
    if not args.soffice or not args.pdftoppm: parser.error("LibreOffice and Poppler are required")
    out=args.output.resolve(); out.mkdir(parents=True, exist_ok=True)
    if any(out.iterdir()): parser.error("Output directory must be empty to prevent stale fixtures")
    docs=out/"docx"; pdfs=out/"pdf"; docs.mkdir(); pdfs.mkdir()
    questions=json.loads((ROOT/"questions.json").read_text())
    records=[]; doc_paths=[]
    for i in range(args.count):
        variant=VARIANTS[i%len(VARIANTS)]; name=f"case_{i+1:04d}_{variant}"
        path=docs/(name+".docx"); expected=make_docx(path,variant,i+1,questions); doc_paths.append(path)
        for suffix in ["docx","pdf"]:
            records.append({"file": f"{suffix}/{name}.{suffix}", "case_id":name, "variant":variant, "scope":"full_form", "expected":expected})
    convert(doc_paths,pdfs,args.soffice)
    with tempfile.TemporaryDirectory(prefix="marp-scan-") as tmp:
        for i,mode in enumerate(["clean", "skewed", "low_resolution", "blurred", "rotated", "mixed_text_scan"]):
            src=pdfs/(doc_paths[i].stem+".pdf"); prefix=Path(tmp)/mode
            subprocess.run([args.pdftoppm,"-r","110" if mode=="low_resolution" else "150","-png",str(src),str(prefix)],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
            pages=sorted(Path(tmp).glob(mode+"-*.png"),key=lambda p:int(p.stem.rsplit("-",1)[1]))
            dest=pdfs/("scan_"+mode+".pdf"); c=canvas.Canvas(str(dest))
            for j,page in enumerate(pages):
                im=Image.open(page).convert("RGB")
                if mode=="skewed": im=im.rotate(2,expand=False,fillcolor="white")
                if mode=="blurred": im=im.filter(ImageFilter.GaussianBlur(1.2))
                if mode=="rotated": im=im.rotate(90,expand=True)
                image_path=Path(tmp)/f"processed-{i}-{j}.png"; im.save(image_path)
                width,height=(612,792) if mode!="rotated" else (792,612)
                c.setPageSize((width,height)); c.drawImage(str(image_path),0,0,width=width,height=height); c.showPage()
            c.save()
            if mode=="mixed_text_scan":
                from pypdf import PdfWriter
                writer=PdfWriter(); native=PdfReader(src); scanned=PdfReader(dest)
                for j,p in enumerate(native.pages): writer.add_page(p if j%2==0 else scanned.pages[j])
                writer.write(dest)
            exp=json.loads(json.dumps(records[2*i]["expected"]))
            # Conservative test policy: scan-derived responses await review.
            for item in exp.values():
                item["review_required"]=True; item["review_reasons"].append("scan_requires_verification")
            records.append({"file":"pdf/"+dest.name,"case_id":dest.stem,"variant":mode,"scope":"full_form","expected":exp})
    for both in [False,True]:
        dest=pdfs/("interactive_both.pdf" if both else "interactive_clear.pdf")
        c=canvas.Canvas(str(dest),pagesize=(612,792)); c.setFont("Helvetica",14)
        c.drawString(40,750,"SYNTHETIC interactive Yes/No probe")
        exp={}
        for i,fid in enumerate(sorted(YESNO)):
            y=680-i*150; c.setFont("Helvetica",10); c.drawString(40,y,fid+" — fictional selection test")
            for option,x in [("Yes",40),("No",180)]:
                c.acroForm.checkbox(name=fid+"_"+option,x=x,y=y-40,checked=both or option=="Yes",buttonStyle="check",size=18)
                c.drawString(x+25,y-35,option)
            exp[fid]={"normalized_answer":None if both else "Yes","selected_options":["Yes","No"] if both else ["Yes"],"review_required":both,"review_reasons":["both_options_selected"] if both else []}
        c.showPage(); c.save()
        fields=PdfReader(dest).get_fields(); assert len(fields)==6
        for key,field in fields.items(): assert str(field["/V"]) == ("/Yes" if both or key.endswith("_Yes") else "/Off")
        records.append({"file":"pdf/"+dest.name,"case_id":dest.stem,"variant":"interactive","scope":"yes_no_probe","expected":exp})
    for mode in ["radio_buttons", "drawn_circles", "offset_ticks", "faint_marks"]:
        dest=pdfs/("selection_"+mode+".pdf")
        c=canvas.Canvas(str(dest),pagesize=(612,792)); c.setFont("Helvetica",14)
        c.drawString(40,750,"SYNTHETIC graphical selection probe")
        exp={}
        for i,fid in enumerate(sorted(YESNO)):
            y=680-i*150; chosen="No" if i==1 else "Yes"
            c.setFillColorRGB(0,0,0); c.setStrokeColorRGB(0,0,0); c.setFont("Helvetica",11)
            c.drawString(40,y,fid+" - fictional selection test")
            for option,x in [("Yes",50),("No",200)]:
                c.drawString(x+25,y-35,option)
                if mode=="radio_buttons":
                    c.acroForm.radio(name=fid,value=option,selected=option==chosen,x=x,y=y-40,size=18,buttonStyle="circle")
                elif mode=="drawn_circles":
                    if option==chosen:
                        c.setLineWidth(1.2); c.ellipse(x+20,y-43,x+60,y-20)
                else:
                    c.setLineWidth(.6); c.rect(x,y-40,16,16)
                    if option==chosen:
                        c.saveState(); c.setLineWidth(1.4)
                        if mode=="faint_marks": c.setStrokeColorRGB(.83,.83,.83)
                        # Mark extends outside the checkbox; it is not a text glyph.
                        c.line(x-3,y-32,x+5,y-42); c.line(x+5,y-42,x+22,y-19); c.restoreState()
            reasons=[] if mode=="radio_buttons" else ["nonstandard_selection_mark"]
            if mode=="faint_marks": reasons=["faint_selection_mark"]
            exp[fid]={"normalized_answer":chosen,"selected_options":[chosen],"review_required":bool(reasons),"review_reasons":reasons}
        c.showPage(); c.save()
        if mode=="radio_buttons":
            fields=PdfReader(dest).get_fields(); assert len(fields)==3
            for fid,expected in exp.items(): assert str(fields[fid]["/V"])=="/"+expected["normalized_answer"]
        records.append({"file":"pdf/"+dest.name,"case_id":dest.stem,"variant":mode,"scope":"yes_no_probe","expected":exp})
    for record in records:
        path=out/record["file"]; assert path.exists(),path
        record["sha256"]=hashlib.sha256(path.read_bytes()).hexdigest()
        if path.suffix==".pdf": record["pages"]=len(PdfReader(path).pages)
    (out/"expected.jsonl").write_text("\n".join(json.dumps(r,ensure_ascii=False) for r in records)+"\n",encoding="utf-8")
    summary={"files":len(records),"docx":args.count,"pdf":args.count+12,"variants":VARIANTS,"synthetic":True}
    (out/"manifest.json").write_text(json.dumps(summary,indent=2)+"\n")
    print(json.dumps(summary))


if __name__ == "__main__": main()
