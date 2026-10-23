"""Local, heuristic questionnaire discovery; no extracted text in the report."""
import argparse
import csv
import io
import re
import shutil
import subprocess
import tempfile
from difflib import SequenceMatcher
from functools import lru_cache
from pathlib import Path

from .extract import read_docx, read_pdf, load_schemas, norm
from .__main__ import ignored_input

MIN_CUES = 8
PROMPT_START = re.compile(
    r"^(?:(?:[a-z]|\d{1,3})[.)]\s*)?(?:what|how|why|when|where|which|who|"
    r"do|does|did|is|are|was|were|have|has|can|will|would|could|please|describe|"
    r"list|provide|indicate|explain|identify|state|quel|quelle|quels|quelles|"
    r"comment|pourquoi|quand|avez|est|veuillez|décrivez|indiquez|précisez)\b", re.I)
IMAGES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}


def cue_key(value):
    value = re.sub(r"^\s*(?:\d{1,3}|[a-z])[.)]\s*", "", value, flags=re.I)
    return norm(value).replace("e mail", "email")


def question_cues(texts):
    """Use known prompts plus bounded question-like spans for unknown forms.

    Cues stay in memory: unknown prompts can accidentally contain response text.
    No cue strings or source snippets are written to the report.
    """
    known = [q["question"] for family in load_schemas().values() for q in family]
    combined = " " + cue_key("\n".join(texts)) + " "
    cues = {cue_key(q) for q in known if " " + cue_key(q) + " " in combined}
    for text in texts:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        for i, line in enumerate(lines):
            if not PROMPT_START.match(line):
                continue
            candidate = line
            for following in lines[i+1:i+9]:
                if re.search(r"[?:.]", candidate) or PROMPT_START.match(following):
                    break
                candidate += " " + following
            # Do not retain inline answers after the question's punctuation.
            stop = re.search(r"[?:.]", candidate)
            if stop is None:
                continue
            candidate = candidate[:stop.start()]
            key = cue_key(candidate)
            if not 4 <= len(key.split()) <= 120:
                continue
            if not any(key == old or key in old or old in key for old in cues):
                cues.add(key)
    return frozenset(cues)


def find_tesseract(explicit=None):
    if explicit:
        path = Path(explicit)
        return str(path) if path.is_file() else shutil.which(explicit)
    found = shutil.which("tesseract")
    if found:
        return found
    for path in [Path("C:/Program Files/Tesseract-OCR/tesseract.exe"),
                 Path("C:/Program Files (x86)/Tesseract-OCR/tesseract.exe")]:
        if path.is_file():
            return str(path)
    return None



class OCRFailure(RuntimeError):
    """Only predefined, content-free diagnostic messages belong here."""


def failure_reason(exc, stage):
    if isinstance(exc, OCRFailure): return str(exc)
    if isinstance(exc, subprocess.TimeoutExpired): return "OCR timed out on a page; try a smaller or clearer PDF"
    if isinstance(exc, (ImportError, ModuleNotFoundError)):
        return "A required Python package could not load; run: python -m pip install -r requirements.txt"
    if isinstance(exc, PermissionError): return "Access denied; check file permissions and local availability"
    if isinstance(exc, FileNotFoundError): return "File or required program not found; ensure the source is downloaded locally"
    if type(exc).__name__ in {"PDFPasswordIncorrect", "FileNotDecryptedError", "WrongPasswordError"}:
        return "Password-protected PDF; provide an unlocked copy in your approved environment"
    if type(exc).__name__ in {"BadZipFile", "PackageNotFoundError"}:
        return "Word document could not be opened; download it locally and confirm it opens in Word"
    return f"{stage} failed ({type(exc).__name__}); check the original opens locally and share this error type for diagnosis"


def check_ocr(executable, language):
    try:
        proc = subprocess.run([executable, "--list-langs"], capture_output=True, timeout=15)
        if proc.returncode: return "Tesseract could not start or read its language data; check the OCR installation"
        installed = set(proc.stdout.decode("utf-8", "replace").splitlines())
        if not set(language.split("+")) <= installed:
            return "Requested OCR language data is missing; install it or select an installed language with --ocr-language"
    except Exception as exc:
        return failure_reason(exc, "OCR startup")
    return ""

def ocr_texts(path, executable, language):
    """Render originals into temporary images; OCR TSV provides ordered lines."""
    import pypdfium2 as pdfium
    from PIL import Image, ImageOps
    texts = []
    pdf = pdfium.PdfDocument(str(path)) if path.suffix.lower() == ".pdf" else None
    image = None if pdf else Image.open(path)
    try:
        count = len(pdf) if pdf else getattr(image, "n_frames", 1)
        with tempfile.TemporaryDirectory(prefix="marp-discovery-") as tmp:
            for index in range(count):
                print(f"  OCR page {index+1}/{count}", flush=True)
                if pdf:
                    page = pdf[index]
                    bitmap = page.render(scale=300/72)
                    frame = bitmap.to_pil().copy()
                    bitmap.close(); page.close()
                else:
                    image.seek(index)
                    frame = ImageOps.exif_transpose(image.copy()).convert("RGB")
                try:
                    image_path = Path(tmp)/"page.png"
                    frame.save(image_path)
                finally:
                    frame.close()
                proc = subprocess.run([executable, str(image_path), "stdout", "-l", language,
                                       "--psm", "3", "tsv"], capture_output=True, timeout=120)
                if proc.returncode:
                    message=proc.stderr.decode("utf-8", "replace").casefold()
                    if "traineddata" in message or "failed loading language" in message:
                        raise OCRFailure("OCR language data could not load; check Tesseract language installation")
                    if "read_params_file" in message and "tsv" in message:
                        raise OCRFailure("Tesseract TSV configuration is missing; repair the OCR installation")
                    raise OCRFailure(f"Tesseract OCR failed (exit code {proc.returncode}); check its installation and source image")
                lines = {}
                for word in csv.DictReader(io.StringIO(proc.stdout.decode("utf-8", "replace")), delimiter="\t"):
                    if word.get("level") == "5" and word.get("text", "").strip():
                        key = tuple(word[k] for k in ["page_num", "block_num", "par_num", "line_num"])
                        lines.setdefault(key, []).append(word["text"])
                texts.append("\n".join(" ".join(words) for words in lines.values()))
    finally:
        if pdf: pdf.close()
        if image: image.close()
    return texts


def inspect_file(path, base, executable, language, ocr_enabled, ocr_issue=""):
    record = dict(submission_id=path.stem, source=str(path.relative_to(base)),
                  status="Ready", method="", reason="", cues=frozenset(), review=False)
    suffix = path.suffix.lower()
    if suffix not in {".pdf", ".docx"} | IMAGES:
        record.update(status="Unsupported format", reason="Convert this file to PDF or DOCX for discovery", review=True)
        return record
    stage="Reading document"
    try:
        scanned = []
        if suffix in IMAGES:
            texts = []; warnings = []; scanned = [1]
        else:
            blocks, scanned, warnings = read_docx(path) if suffix == ".docx" else read_pdf(path)
            texts = [c.text for b in blocks for c in b.cells]
            record["method"] = "Word text" if suffix == ".docx" else "PDF text"
        stage="Identifying question cues"
        cues = question_cues(texts)
        record["cues"] = cues
        needs_ocr = bool(scanned) or len(cues) < MIN_CUES
        if needs_ocr and suffix != ".docx":
            if ocr_issue:
                record.update(status="OCR unavailable", reason=ocr_issue, review=True)
                return record
            if executable and ocr_enabled:
                # Keep usable native cues too, but mark the grouping as OCR-dependent.
                try:
                    cues |= question_cues(ocr_texts(path, executable, language))
                except Exception as exc:
                    record.update(status="OCR failed", reason=failure_reason(exc,"OCR reading")+"; native question cues retained, version left unassigned", review=True)
                    return record
                record["method"] = "Local OCR" if not texts else "PDF text + local OCR"
                record["review"] = True
                record["reason"] = "OCR-assisted grouping; compare questions in the original"
            elif scanned or not " ".join(texts).strip():
                record.update(status="Needs OCR", reason="Install local Tesseract OCR, then rerun discovery", review=True)
                return record
        if warnings:
            record["review"] = True
            record["reason"] += "; Reader detected content needing manual review"
        record["cues"] = cues
        if len(cues) < MIN_CUES:
            record.update(status="Not enough question evidence", review=True,
                          reason="Open manually; too few recognizable questions to assign a version")
    except Exception as exc:
        # Do not export parser messages: they can contain source text.
        record.update(status="Read failed", reason=failure_reason(exc,stage), review=True)
    return record


@lru_cache(maxsize=8192)
def similarity(left, right):
    common = len(left & right)
    a = sorted(left - right); b = set(right - left)
    for value in a:
        matches = [(SequenceMatcher(None, value, other).ratio(), other) for other in b
                   if abs(len(value)-len(other)) <= max(len(value),len(other))*.08]
        if matches:
            score, other = max(matches)
            if score >= .97:
                common += 1; b.remove(other)
    union = len(left)+len(right)-common
    return common/union if union else 0.0


def group_records(records, threshold):
    groups = []
    for record in records:
        record["group"] = ""; record["representative"] = ""; record["similarity"] = ""
        if record["status"] != "Ready": continue
        matches = []
        for index, group in enumerate(groups):
            # Compare unique signatures: all members must agree, avoiding chains
            # of partial forms that bridge two distinct versions.
            score = min(similarity(record["cues"], cues) for cues in group["signatures"])
            if score >= threshold: matches.append((score, index))
        if matches:
            _, index = max(matches)
            group = groups[index]
            group["members"].append(record);group["signatures"].add(record["cues"])
        else:
            groups.append(dict(members=[record],signatures={record["cues"]}))
    for index, group in enumerate(groups, 1):
        representative = sorted(group["members"], key=lambda r: (-len(r["cues"]), "OCR" in r["method"], r["source"]))[0]
        for record in group["members"]:
            record["group"] = f"Candidate {index:03d}"
            record["representative"] = representative["source"]
            record["similarity"] = f"{similarity(record['cues'],representative['cues']):.0%}"
    return groups


def save_report(records, target, threshold):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    groups = group_records(records, threshold)
    wb = Workbook();ws = wb.active;ws.title = "Version discovery"
    headers = ["Submission ID", "Flagged", "Candidate version", "Representative file to open",
               "Use this file as representative?", "Files in candidate group", "Status", "Reason",
               "Question cues found", "Similarity to representative", "Read method", "Source file"]
    ws.append(headers)
    sizes = {g["members"][0]["group"]:len(g["members"]) for g in groups}
    for r in records:
        representative = bool(r["group"]) and r["source"] == r["representative"]
        reason = r["reason"].strip("; ")
        if representative:
            reason = (reason+"; " if reason else "")+"Review this candidate version; grouping is not confirmation"
        row = [r["submission_id"], "Yes" if r["review"] or representative else "No", r["group"],
               r["representative"], "Yes" if representative else "No", sizes.get(r["group"], ""),
               r["status"],reason,len(r["cues"]),r["similarity"],r["method"],r["source"]]
        ws.append([ILLEGAL_CHARACTERS_RE.sub("?",v)[:32767] if isinstance(v,str) else v for v in row])
    for row in ws:
        for c in row:
            if isinstance(c.value,str):c.data_type="s"
            c.alignment=Alignment(wrap_text=True,vertical="top");c.font=Font(name="Arial",size=11)
    for c in ws[1]:
        c.font=Font(name="Arial",bold=True,color="FFFFFF",size=11)
        c.fill=PatternFill("solid",fgColor="243C59")
    ws.row_dimensions[1].height=45
    for i in range(2,ws.max_row+1):
        ws.row_dimensions[i].height=60
        if ws.cell(i,2).value=="Yes":ws.cell(i,2).fill=PatternFill("solid",fgColor="FFF0C2")
    for i,width in enumerate([20,12,22,48,24,18,30,65,18,22,24,48],1):
        ws.column_dimensions[get_column_letter(i)].width=width
    ws.freeze_panes="C2";ws.auto_filter.ref=ws.dimensions
    target=Path(target);temp=target.with_suffix(".tmp.xlsx");wb.save(temp);temp.replace(target)
    return len(groups)


def run(input_path, output, threshold=.96, ocr="auto", tesseract=None, language="eng"):
    source=Path(input_path).resolve();output=Path(output).resolve()
    if not source.exists():raise ValueError("Input path does not exist")
    if source==output or output in source.parents:raise ValueError("Output must not equal or contain the input directory")
    if not 0.5<=threshold<=1:raise ValueError("Similarity must be between 0.5 and 1")
    base=source.parent if source.is_file() else source
    paths=[source] if source.is_file() else sorted(p for p in source.rglob("*") if p.is_file() and output not in p.parents)
    if not paths:raise ValueError("No files found")
    output.mkdir(parents=True,exist_ok=True);target=output/"questionnaire_versions.xlsx"
    executable=find_tesseract(tesseract) if ocr!="off" else None
    if tesseract and not executable:raise ValueError("Tesseract executable not found at the specified path")
    ocr_issue=check_ocr(executable,language) if executable else ""
    print("Local OCR available" if executable and not ocr_issue else "Local OCR unavailable/off: "+(ocr_issue or "unreadable scans will be marked Needs OCR"),flush=True)
    records=[];interrupted=False
    try:
        for index,path in enumerate(paths,1):
            print(f"[{index}/{len(paths)}] Inspecting {path.relative_to(base)}",flush=True)
            if ignored_input(path,base):
                r=dict(submission_id=path.stem,source=str(path.relative_to(base)),status="Skipped",method="",reason="Temporary or hidden file",cues=frozenset(),review=False)
            else:r=inspect_file(path,base,executable,language,ocr!="off",ocr_issue)
            records.append(r)
            print(f"  {r['status']}; {len(r['cues'])} question cues",flush=True)
            if r["status"] not in {"Ready","Skipped"}: print("  Reason: "+r["reason"],flush=True)
            if index%25==0:
                save_report(records,target,threshold)
                print(f"  Progress saved: {index} files",flush=True)
    except KeyboardInterrupt:
        interrupted=True
    count=save_report(records,target,threshold)
    print(f"{'Stopped early. ' if interrupted else ''}Saved {target}: {len(records)} files inspected, {count} candidate groups. Review representatives; no answers or question snippets exported.",flush=True)
    return records


def main():
    parser=argparse.ArgumentParser(description="Discover candidate questionnaire versions locally; export a file inventory to Excel.")
    parser.add_argument("--input",required=True,type=Path);parser.add_argument("--output",required=True,type=Path)
    parser.add_argument("--similarity",type=float,default=.96,help="Grouping threshold, 0.5 to 1 (default .96)")
    parser.add_argument("--ocr",choices=["auto","off"],default="auto")
    parser.add_argument("--tesseract",help="Optional full path to tesseract.exe")
    parser.add_argument("--ocr-language",default="eng",help="Installed Tesseract language, e.g. eng+fra")
    args=parser.parse_args()
    try:run(args.input,args.output,args.similarity,args.ocr,args.tesseract,args.ocr_language)
    except (ValueError,OSError) as exc:parser.exit(2,f"Error: {exc}\n")


if __name__=="__main__":main()
