"""Regression checks using fictional responses, with no user documents."""
import copy
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from marp.extract import Block, Cell, extract, yes_no, load_schemas
from marp.__main__ import run


def row(*values):
    return Block([Cell(x,x) for x in values], {'table':1}, True)


def parsed(blocks):
    return extract(Path('001.pdf'),reader=lambda _: (copy.deepcopy(blocks),[],[]))


class PipelineTests(unittest.TestCase):
    def test_contacts_and_separate_questions(self):
        r=parsed([row('Events Component'),row('Name of your organization:', 'Example Group'),
                  row('Project Contact name'),row('First name','Last name'),row('Alex','Example'),
                  row('Email address:','alex@example.invalid'),
                  row('Event start date:', '2026-06-12'),row('Event end date:', '2026-06-13')])
        a={a['field_id']:a['answer_text'] for a in r['responses']}
        self.assertEqual([a[k] for k in ['A01','A02','A03','A07','B02','B03']],
                         ['Example Group','Alex','Example','alex@example.invalid','2026-06-12','2026-06-13'])

    def test_yes_no_never_guesses_blank_or_both(self):
        for value in ['[X] Yes [X] No', 'Yes / No', '☐ Oui ☐ Non']:
            answer,reasons=yes_no(value)
            self.assertIsNone(answer)
            self.assertTrue(reasons)
        self.assertEqual(yes_no('Oui ☒ / Non ☐')[0],'Yes')
        self.assertEqual(yes_no('[X] No [ ] Yes')[0],'No')

    def test_repeated_explain_prompts_stay_with_parent(self):
        fields=load_schemas()['marp_events_legacy']
        blocks=[row('Expected total number of participants','100')]
        for i in range(1,5):
            q=next(q for q in fields if q['field_id']==f'OLD-D0{i}')
            blocks += [row(q['question'],'Yes'),row('Explain how:',f'Fictional result {i}')]
        r=parsed(blocks)
        a={a['field_id']:a for a in r['responses']}
        for i in range(1,5):
            self.assertEqual(a[f'OLD-D0{i}']['normalized_answer'],'Yes')
            self.assertEqual(a[f'OLD-D0{i}E']['answer_text'],f'Fictional result {i}')

    def test_french_contacts_and_choices(self):
        r=parsed([row('Nom de votre organisme','Groupe fictif'),row('Prénom','Émilie'),
                  row('Nom de famille','Exemple'),row('Adresse électronique','emilie@example.invalid'),
                  row('Votre événement a-t-il changé par rapport à ce qui était initialement prévu','Oui ☒ / Non ☐'),
                  row('Dans l’affirmative, veuillez expliquer les principaux changements qui ont eu lieu et les raisons de ces derniers.','Lieu modifié.')])
        self.assertEqual(r['template_id'],'marp_events_fr')
        a={a['field_id']:a for a in r['responses']}
        self.assertEqual(a['FR-A07']['answer_text'],'emilie@example.invalid')
        self.assertEqual(a['FR-C02']['normalized_answer'],'Yes')
        self.assertEqual(a['FR-C03']['answer_text'],'Lieu modifié.')

    def test_french_older_version_not_misclassified(self):
        r=parsed([row('Nom de votre organisme'),row('Nombre total prévu de participants')])
        self.assertEqual(r['template_id'],'unknown')
        self.assertIn('unsupported_french_legacy_template',r['review_reasons'])

    def test_respondent_note_preserved_known_note_removed(self):
        from marp.extract import PRINTED_NOTES
        r=parsed([row('Events Component'),row('4.1 What helped make your event a success?',
                  'Note: volunteers stayed late.\n'+PRINTED_NOTES[0]),row('Name and signature (required)')])
        a=next(a for a in r['responses'] if a['field_id']=='D08')
        self.assertEqual(a['answer_text'],'Note: volunteers stayed late.')

    def test_scans_and_mixed_pdfs_skipped_without_answers(self):
        blocks=[row("Events Component"),row("First name","Alex")]
        r=extract(Path("001.pdf"),reader=lambda _: (blocks,[2],[]))
        self.assertEqual(r["status"],"skipped_ocr_required")
        self.assertEqual(r["responses"],[])
        self.assertEqual(r["ocr_pages"],[2])
        self.assertTrue(r["review_required"])

    def test_image_only_word_skipped(self):
        r=extract(Path("001.docx"),reader=lambda _: ([],[],["embedded_images_require_review"]))
        self.assertEqual(r["status"],"skipped_ocr_required")

    def test_unknown_form_text_preserved_for_review(self):
        r=parsed([row("Unrecognized questionnaire", "Fictional answer")])
        self.assertEqual(r["status"],"unsupported_template")
        self.assertIn("Fictional answer",r["unmatched"][0]["text"])

    def test_valid_yes_no_export_is_plain_answer(self):
        from marp.export import workbook
        r=parsed([row("Events Component"),row("Has your event changed from what was originally planned (refer to your funding application)?","[X] Yes [ ] No")])
        r.update(submission_id="001",relative_path="001.pdf")
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/"answers.xlsx";workbook([r],p)
            ws=load_workbook(p).active
            col=next(c.column for c in ws[1] if str(c.value).startswith("Has your event changed"))
            self.assertEqual(ws.cell(2,col).value,"Yes")

    def test_batch_failure_and_ids_excel(self):
        with tempfile.TemporaryDirectory() as td:
            base=Path(td);src=base/'forms';src.mkdir()
            for name in ['001.pdf','002.pdf','003.doc']:(src/name).touch()
            def reader(path):
                if path.stem=='002':raise ValueError('Do not expose source text')
                return [row('Events Component'),row('Name of your organization:','=1+1')],[],[]
            out=base/'results'
            results=run(src,out,reader=reader)
            ws=load_workbook(out/'extracted_responses.xlsx').active
            self.assertEqual(ws.cell(2,1).value,'001')
            self.assertEqual(ws.cell(1,2).value,'Flagged')
            self.assertEqual(ws.cell(2,3).value,'=1+1')
            self.assertEqual(ws.cell(2,3).data_type,'s')
            self.assertEqual([r['status'] for r in results],['extracted','error','unsupported_format'])
            self.assertNotIn('Do not expose source text',str(list(ws.values)))
            self.assertFalse(any('Authorized Signature' in str(x.value) for x in ws[1]))

    def test_completed_rows_saved_before_interruption(self):
        with tempfile.TemporaryDirectory() as td:
            base=Path(td);src=base/'forms';src.mkdir()
            for name in ['001.pdf','002.pdf']:(src/name).touch()
            def reader(path):
                if path.stem=='002':raise KeyboardInterrupt()
                return [row('Events Component')],[],[]
            with self.assertRaises(KeyboardInterrupt):run(src,base/'results',reader=reader)
            ws=load_workbook(base/'results/extracted_responses.xlsx').active
            self.assertEqual(ws.max_row,2)
            self.assertEqual(ws.cell(2,1).value,'001')


if __name__=='__main__':unittest.main()
