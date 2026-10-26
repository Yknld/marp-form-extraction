"""Worksheets by questionnaire version, with one source file per row."""
from pathlib import Path
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from .extract import load_schemas


def text(value):
    if value is None:
        return ""
    return ILLEGAL_CHARACTERS_RE.sub(lambda m: "\\u%04x" % ord(m.group()), str(value))


def readable(reason):
    return reason.replace("_", " ")


SHEET_FAMILIES = {
    "English Type 1": "marp_events",
    "English Type 2": "marp_events_legacy",
    "French Type 1": "marp_events_fr",
    "French Type 2": None,  # Recognizable, but no verified extraction schema yet.
}


def sheet_for(result):
    for sheet, family in SHEET_FAMILIES.items():
        if family and result.get("template_id") == family:
            return sheet
    if "unsupported_french_legacy_template" in result.get("review_reasons", []):
        return "French Type 2"
    return "Miscellaneous"


def workbook(results, target):
    schemas = load_schemas()
    groups = {name: [] for name in [*SHEET_FAMILIES, "Miscellaneous"]}
    for result in results:
        groups[sheet_for(result)].append(result)
    wb = Workbook()
    wb.remove(wb.active)
    for name, group in groups.items():
        family = SHEET_FAMILIES.get(name)
        families = [family] if family else [
            key for key in schemas if any(r.get("template_id") == key for r in group)]
        write_sheet(wb.create_sheet(name), group, schemas, families)
    target = Path(target)
    tmp = target.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(target)


def write_sheet(ws, results, schemas, families):
    fields = [(family, q) for family in families for q in schemas[family] if q["response_type"] != "Signature"]
    counts = Counter(q["question"].casefold() for _, q in fields)
    labels = {"marp_events": "Events", "marp_events_legacy": "Events older version", "marp_events_fr": "Événements français", "marp_anti_racism_call": "Anti-Racism", "marp_organizational_capacity": "Capacity Building"}
    question_types={(family,q["field_id"]):q["response_type"] for family,q in fields}
    columns = []
    for family, q in fields:
        label = q["question"]
        if counts[label.casefold()] > 1:
            label += f" [{labels[family]} — {q['field_id']}]"
        columns.append(((family, q["field_id"]), label))
    # Preserve each activity's correspondence across the three cells.
    activity_columns = [("description", "Activities — description"),
                        ("completed_raw", "Activities — completed?"),
                        ("participant_count_raw", "Activities — participant count")]
    if any(r["template_id"]=="marp_organizational_capacity" for r in results):
        activity_columns.insert(2,("results_achieved","Activities — results achieved"))
    has_activities = any(r.get("activities") for r in results)
    headers = ["Submission ID", "Flagged"] + [label for _, label in columns]
    if has_activities:
        headers += [label for _, label in activity_columns]
    headers += ["Reason", "Source file", "Status", "Unmatched text"]
    rows = []
    for result in results:
        answers = {a["field_id"]: a for a in result.get("responses", [])}
        reasons = [readable(x) for x in result.get("review_reasons", [])]
        row = [text(result["submission_id"]), "No"]
        for (family, fid), label in columns:
            a = answers.get(fid) if result["template_id"] == family else None
            value=a["answer_text"] if a else ""
            if a and question_types[(family,fid)]=="Yes / No" and a.get("normalized_answer") in {"Yes","No"}:
                value=a["normalized_answer"]
            row.append(text(value))
            if a:
                reasons.extend(f"{label}: {readable(x)}" for x in a.get("review_reasons", []))
        if has_activities:
            for key, _ in activity_columns:
                row.append("\n\n".join(f"Activity {a['activity_index']}: {text(a.get(key))}" for a in result.get("activities", [])))
        for a in result.get("activities", []):
            reasons.extend(f"Activity {a['activity_index']}: {readable(x)}" for x in a.get("review_reasons", []))
        for a in result.get("unmatched", []):
            reasons.append("Unmatched text: " + readable(a["review_reason"]))
        if result.get("error"):
            reasons.append(result["error"])
        flagged = result.get("review_required", False) or bool(reasons)
        row[1] = "Yes" if flagged else "No"
        row += ["\n".join(dict.fromkeys(reasons)),
                text(result["relative_path"]), readable(result["status"]).capitalize(),
                "\n\n".join(a["text"] for a in result.get("unmatched", []))]
        rows.append([text(v) for v in row])
    # Excel cells hold at most 32,767 characters. Continue across columns,
    # keeping one row per file, without losing text.
    expanded_headers = []
    expanded_rows = [[] for _ in rows]
    for col, header in enumerate(headers):
        parts = max([1] + [(len(row[col]) + 29999) // 30000 for row in rows])
        for part in range(parts):
            expanded_headers.append(header if part == 0 else f"{header} (continued {part + 1})")
            for source, dest in zip(rows, expanded_rows):
                dest.append(source[col][part * 30000:(part + 1) * 30000])
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "C2"
    ws.append(expanded_headers)
    for row in expanded_rows:
        ws.append(row)
    for row in ws:
        for cell in row:
            cell.data_type = "s"  # Preserve IDs and prevent response formulas.
            cell.font = Font(name="Arial", size=11)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243C59")
    ws.row_dimensions[1].height = 120
    for i, header in enumerate(expanded_headers, 1):
        width = 48
        if header == "Submission ID": width = 20
        elif header == "Flagged": width = 12
        elif header == "Status": width = 24
        ws.column_dimensions[get_column_letter(i)].width = width
    for i in range(2, ws.max_row + 1):
        ws.row_dimensions[i].height = 90
        if i % 2 == 0:
            for cell in ws[i]: cell.fill = PatternFill("solid", fgColor="EEF3F8")
    ws.auto_filter.ref = ws.dimensions
    if rows:
        flag = get_column_letter(expanded_headers.index("Flagged") + 1)
        validation = DataValidation(type="list", formula1='"Yes,No"')
        ws.add_data_validation(validation)
        validation.add(f"{flag}2:{flag}{ws.max_row}")
        ws.conditional_formatting.add(f"{flag}2:{flag}{ws.max_row}", FormulaRule(
            formula=[f'{flag}2="Yes"'], fill=PatternFill("solid", fgColor="FFF0C2")))
