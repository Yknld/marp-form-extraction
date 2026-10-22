import contextlib
import io
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from docx import Document
from openpyxl import load_workbook
from marp.extract import yes_no,normalize_value,extract
from marp.__main__ import run
from marp.export import workbook

ROOT=Path(__file__).resolve().parents[1]


class SelectionTests(unittest.TestCase):
    def test_clear_selection_positions(self):
        for text,expected in [("Yes ☑ / No ☐","Yes"),("☐ Yes ☒ No","No"),("[X] No [ ] Yes","No"),("Oui ✓ / Non ☐","Yes"),("Yes X / No","Yes"),("No","No")]:
            self.assertEqual(yes_no(text)[0],expected,text)

    def test_ambiguity_is_not_yes(self):
        for text in ["Yes X / No X","[X] Yes [X] No","Yes ☐ / No ☐","Yes / No"]:
            value,reasons=yes_no(text);self.assertIsNone(value,text);self.assertTrue(reasons,text)

    def test_dates_and_zero(self):
        self.assertEqual(normalize_value("0","Number"),(0,[]))
        self.assertIsNone(normalize_value("","Number")[0])
        self.assertIsNone(normalize_value("N/A","Number")[0])
        self.assertEqual(normalize_value("2026-02-30","Date")[1],["invalid_date"])
        self.assertEqual(normalize_value("06/07/2026","Date")[1],["ambiguous_date"])


class ExtractionTests(unittest.TestCase):
    def test_split_activity_and_spacer(self):
        for name in ["arc_03_split_activity_row.docx","arc_06_blank_spacer_rows.docx"]:
            result=extract(ROOT/"fixtures/anti_racism_call/docx"/name)
            self.assertEqual(len(result["activities"]),4)
            self.assertEqual(result["activities"][1]["participant_count"],20)
            self.assertIn("shared examples",result["activities"][1]["description"])

    def test_missing_question_and_header_words_in_answer(self):
        result=extract(ROOT/"fixtures/synthetic/docx/case_0013_missing_question.docx")
        item=next(a for a in result["responses"] if a["field_id"]=="D07")
        self.assertEqual(item["presence"],"missing");self.assertTrue(item["review_required"])
        result=extract(ROOT/"fixtures/synthetic/docx/case_0017_header_words_in_answer.docx")
        item=next(a for a in result["responses"] if a["field_id"]=="D09")
        self.assertIn('"E. Partnerships"',item["answer_text"])

    def test_scan_routes_to_ocr(self):
        result=extract(ROOT/"fixtures/synthetic/pdf/scan_clean.pdf")
        self.assertEqual(result["status"],"ocr_required");self.assertTrue(result["ocr_pages"])

    def test_wrapped_section_heading_is_not_answer(self):
        r=extract(ROOT/"fixtures/synthetic/pdf/case_0001_ascii_x.pdf")
        a=next(a for a in r["responses"] if a["field_id"]=="F01")
        self.assertNotIn("Acknowledgement",a["answer_text"])

    def test_unsupported_language_is_flagged(self):
        from marp.extract import Block,Cell,choose_template,load_schemas
        template,language,reasons=choose_template([Block([Cell("Nom de votre organisme")],{})],load_schemas())
        self.assertEqual(template,"unknown");self.assertEqual(language,"fr");self.assertIn("unsupported_french_template",reasons)

    def test_blank_optional_field(self):
        r=extract(ROOT/"fixtures/synthetic/docx/case_0015_numeric_edge_cases.docx")
        a=next(a for a in r["responses"] if a["field_id"]=="A06")
        self.assertEqual(a["presence"],"blank");self.assertFalse(a["review_required"])

    def test_batch_resume_error_duplicate_and_changes(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);src=root/"input";src.mkdir();out=root/"output"
            original=ROOT/"fixtures/synthetic/docx/case_0001_ascii_x.docx"
            shutil.copy(original,src/"one.docx");shutil.copy(original,src/"copy.docx")
            (src/"broken.pdf").write_bytes(b"not a pdf")
            with contextlib.redirect_stdout(io.StringIO()):
                results=run(src,out);run(src,out)
            self.assertEqual(sum(r["status"]=="error" for r in results),1)
            self.assertEqual(sum("duplicate_of" in r for r in results),1)
            self.assertEqual(json.loads((out/"summary.json").read_text())["resumed"],3)
            doc=Document(src/"one.docx");doc.add_paragraph("Changed input");doc.save(src/"one.docx")
            with contextlib.redirect_stdout(io.StringIO()):run(src,out)
            self.assertEqual(json.loads((out/"summary.json").read_text())["resumed"],2)
            self.assertTrue((out/"extracted_responses.xlsx").exists())

    def test_excel_formula_text_remains_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            r=extract(ROOT/"fixtures/synthetic/docx/case_0001_ascii_x.docx")
            r.update(file_id="x",relative_path="=1+1.docx")
            r["responses"][0]["answer_text"]="=HYPERLINK(\"https://example.invalid\")"
            path=Path(tmp)/"test.xlsx";workbook([r],path);w=load_workbook(path)
            self.assertEqual(w["Responses"]["F2"].data_type,"s")
            self.assertEqual(w["Files"]["B2"].data_type,"s")

    def test_mixed_folder_tracks_every_format_and_submission_id(self):
        from pypdf import PdfWriter
        with tempfile.TemporaryDirectory() as tmp:
            src=Path(tmp)/"forms";src.mkdir();nested=src/"batch";nested.mkdir();out=src/"results"
            shutil.copy(ROOT/"fixtures/synthetic/docx/case_0001_ascii_x.docx",src/"001234.DOCX")
            shutil.copy(ROOT/"fixtures/synthetic/pdf/case_0001_ascii_x.pdf",nested/"001234.PDF")
            for name in ["old.doc","image.png","bundle.zip","notes.txt","unknown","broken.docx"]:
                (src/name).write_bytes(b"test input")
            writer=PdfWriter();writer.add_blank_page(width=100,height=100);writer.encrypt("secret")
            with (src/"locked.pdf").open("wb") as f:writer.write(f)
            for name in ["~$001234.docx",".DS_Store","draft.tmp","Thumbs.db"]:(src/name).write_bytes(b"ignored")
            with contextlib.redirect_stdout(io.StringIO()):
                first=run(src,out);run(src,out)
            self.assertEqual(len(first),9)
            by_name={r["relative_path"]:r for r in first}
            for name in ["001234.DOCX",str(Path("batch")/"001234.PDF")]:
                self.assertEqual(by_name[name]["status"],"extracted")
                self.assertEqual(by_name[name]["submission_id"],"001234")
                self.assertIn("duplicate_submission_id",by_name[name]["review_reasons"])
            for name in ["old.doc","image.png","bundle.zip","notes.txt","unknown"]:
                self.assertEqual(by_name[name]["status"],"unsupported_format")
                self.assertTrue(by_name[name]["review_required"])
            for name in ["broken.docx","locked.pdf"]:self.assertEqual(by_name[name]["status"],"error")
            summary=json.loads((out/"summary.json").read_text())
            self.assertEqual(summary["resumed"],9);self.assertEqual(len(summary["ignored_files"]),4)
            w=load_workbook(out/"extracted_responses.xlsx")
            for sheet in ["Files","Responses"]:
                ws=w[sheet];ids=[row[-1] for row in ws.iter_rows(min_row=2) if row[1].value=="001234.DOCX"]
                self.assertTrue(ids)
                self.assertTrue(all(c.value=="001234" and c.data_type=="s" for c in ids))
            (nested/"001234.PDF").unlink()
            with contextlib.redirect_stdout(io.StringIO()):updated=run(src,out)
            self.assertNotIn("duplicate_submission_id",next(r for r in updated if r["relative_path"]=="001234.DOCX")["review_reasons"])

    def test_unsupported_only_folder_is_reported(self):
        with tempfile.TemporaryDirectory() as tmp:
            src=Path(tmp)/"forms";src.mkdir();(src/"001.odt").write_bytes(b"unsupported")
            with contextlib.redirect_stdout(io.StringIO()):results=run(src,Path(tmp)/"out")
            self.assertEqual(results[0]["submission_id"],"001")
            self.assertEqual(results[0]["status"],"unsupported_format")


if __name__=="__main__":unittest.main()
